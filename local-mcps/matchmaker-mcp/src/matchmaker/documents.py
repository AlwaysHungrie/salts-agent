"""Files to plain text: PDF (text layer), DOCX, TXT and MD. No model call; a scanned PDF with no text layer is
refused so the agent can send the text it read itself."""

from pathlib import Path

from .errors import ToolFailure

SUFFIXES = (".pdf", ".docx", ".txt", ".md")


def read_text(path: Path) -> str:
    if not path.is_file():
        raise ToolFailure("not_found", "no file at that path")
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise ToolFailure("invalid_input", f"unsupported file type {suffix or '(none)'}; send {', '.join(SUFFIXES)}")
    try:
        if suffix == ".pdf":
            text = _pdf(path)
        elif suffix == ".docx":
            text = _docx(path)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
    except ToolFailure:
        raise
    except Exception:
        raise ToolFailure("invalid_input", f"the {suffix} file could not be read") from None
    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not text:
        raise ToolFailure("invalid_input", "the file has no text (a scanned PDF?); send the text instead")
    return text


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    return "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)

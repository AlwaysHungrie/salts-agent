import json

import pytest

from matchmaker import billing
from matchmaker.documents import read_text
from matchmaker.errors import ToolFailure
from matchmaker.jobs import find_candidates, post_job
from matchmaker.judge import check_score, judge_page, score, search_score
from matchmaker.models import Profile
from matchmaker.people import add_document, delete_person, get_person, render_profile
from matchmaker.search import duplicate_groups, fuse, search

from .conftest import resume

# ------------------------------------------------------------------ files to text


def test_reads_pdf_docx_and_text(tmp_path):
    import docx
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(text="Backend engineer, Python and Kafka")
    pdf.output(str(tmp_path / "cv.pdf"))
    assert "Python and Kafka" in read_text(tmp_path / "cv.pdf")

    d = docx.Document()
    d.add_paragraph("Accountant, GST and Tally")
    d.add_table(rows=1, cols=2).rows[0].cells[0].text = "SAP FICO"
    d.save(str(tmp_path / "cv.docx"))
    text = read_text(tmp_path / "cv.docx")
    assert "GST and Tally" in text and "SAP FICO" in text

    (tmp_path / "cv.md").write_text("  Data scientist  \n", encoding="utf-8")
    assert read_text(tmp_path / "cv.md") == "Data scientist"


def test_unreadable_files_are_refused(tmp_path):
    from fpdf import FPDF

    (tmp_path / "cv.exe").write_bytes(b"MZ")
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    pdf = FPDF()
    pdf.add_page()
    pdf.output(str(tmp_path / "scan.pdf"))  # a page with no text layer
    for name, code in [("cv.exe", "invalid_input"), ("broken.pdf", "invalid_input"), ("scan.pdf", "invalid_input"),
                       ("missing.pdf", "not_found")]:
        with pytest.raises(ToolFailure) as e:
            read_text(tmp_path / name)
        assert e.value.code == code, name


# ------------------------------------------------------------------ people


async def test_adding_a_resume_makes_the_folder_and_keeps_contacts_out_of_the_profile(svc, tmp_path):
    src = tmp_path / "Asha CV.md"
    src.write_text(resume("Asha Rao", "Backend engineer", "Python, Django", relocate="no", years=6.5),
                   encoding="utf-8")
    r = await add_document(svc, "resume", file_path=src)
    assert r.status == "created" and r.name == "Asha Rao" and r.headline == "Backend engineer"
    folder = svc.settings.data_dir / "people" / r.person_id
    assert sorted(p.name for p in folder.iterdir()) == ["embedding.npy", "person.json", "profile.md", "resume.md",
                                                        "resume.txt"]
    profile = (folder / "profile.md").read_text(encoding="utf-8")
    assert "- Open to relocate: no" in profile and "- Experience: 6.5 years" in profile
    assert "- Skills: Python, Django" in profile
    assert "asha@example.com" not in profile and "+91" not in profile and "Asha" not in profile
    meta = json.loads((folder / "person.json").read_text(encoding="utf-8"))
    assert meta["email"] == "asha@example.com" and meta["documents"]["resume"]["original"] == "resume.md"
    assert svc.llm.calls == ["person_profile"] and svc.embedder.calls == 1


async def test_a_cover_letter_rewrites_the_profile_from_both_documents(svc):
    r = await add_document(svc, text=resume("Ravi K", "Data analyst", "SQL"))
    r2 = await add_document(svc, "cover_letter", text="Headline: Data scientist\nSkill: PyTorch\nRelocate: yes",
                            person_id=r.person_id)
    assert r2.status == "updated" and r2.person_id == r.person_id and r2.headline == "Data scientist"
    prompt = svc.llm.prompts[-1]
    assert '<document kind="resume">' in prompt and '<document kind="cover_letter">' in prompt
    p = get_person(svc, r.person_id)
    assert p.documents == ["resume", "cover_letter"] and "SQL, PyTorch" in p.profile
    assert "- Open to relocate: yes" in p.profile


async def test_a_new_resume_replaces_the_old_one_and_the_same_text_is_no_model_call(svc, tmp_path):
    first = tmp_path / "cv.md"
    first.write_text(resume("Meera S", "QA engineer", "Selenium"), encoding="utf-8")
    r = await add_document(svc, file_path=first)
    same = await add_document(svc, text=resume("Meera S", "QA engineer", "Selenium"), person_id=r.person_id)
    assert same.status == "unchanged" and len(svc.llm.calls) == 1
    await add_document(svc, text=resume("Meera S", "SDET", "Playwright"), person_id=r.person_id)
    folder = svc.settings.data_dir / "people" / r.person_id
    assert not (folder / "resume.md").exists()  # the replaced original goes
    assert "Playwright" in (folder / "resume.txt").read_text(encoding="utf-8")


async def test_bad_input_is_refused_and_nothing_is_written(svc):
    for kwargs, code in [({"text": "  "}, "invalid_input"), ({"kind": "notes", "text": "x"}, "invalid_input"),
                         ({"text": "x", "person_id": "../../etc"}, "invalid_input"),
                         ({"text": "x", "person_id": "0" * 32}, "not_found")]:
        with pytest.raises(ToolFailure) as e:
            await add_document(svc, **kwargs)
        assert e.value.code == code, kwargs
    assert not (svc.settings.data_dir / "people").exists()


async def test_a_failed_model_call_leaves_no_person(svc):
    async def boom(*a, **k):
        raise ToolFailure("llm_unavailable", "down")
    svc.embedder.embed = boom
    with pytest.raises(ToolFailure):
        await add_document(svc, text=resume("A B", "Engineer", "Go"))
    assert not (svc.settings.data_dir / "people").exists()


async def test_delete_person(svc):
    r = await add_document(svc, text=resume("Tom L", "Designer", "Figma"))
    delete_person(svc, r.person_id)
    with pytest.raises(ToolFailure) as e:
        get_person(svc, r.person_id)
    assert e.value.code == "not_found"


def test_profile_shows_relocation_reason_and_skips_unknowns():
    p = Profile(name="X", email=None, phone=None, headline="Nurse", current_title=None, seniority="senior",
                location=None, open_to_relocate="likely", relocation_note="moved for work before", work_mode=None,
                years_experience=None, work_authorization="likely UK, UK training and roles", notice_period=None,
                languages=[], education=None, industries=["Healthcare"], skills=["ICU care"], summary="ICU nurse.")
    md = render_profile(p)
    assert md.startswith("# Nurse\n")
    assert "- Open to relocate: likely (moved for work before)" in md
    assert "- Location: unknown" in md and "- Experience: unknown" in md
    assert "Notice period" not in md and "Languages" not in md and "X" not in md


def test_profile_prompt_reads_relocation_generously():
    from matchmaker.prompts import render

    system, _ = render("profile", documents="")
    assert '"no" only when they say they will not relocate' in system
    assert "leaning \"likely\"" in system


# ------------------------------------------------------------------ search


def test_fuse_rewards_agreement_and_keeps_everyone():
    order = fuse(keyword=[2, 0], meaning=[0, 1, 2], k=60)
    assert order[0] == 0 and set(order) == {0, 1, 2}
    assert fuse(keyword=[], meaning=[1, 0], k=60) == [1, 0]  # no query word anywhere: meaning alone


async def test_hybrid_ranks_words_and_meaning(svc):
    await add_document(svc, text=resume("Asha R", "Backend engineer", "Python, Django, Kafka"))
    await add_document(svc, text=resume("Bo C", "Frontend engineer", "React, TypeScript"))
    await add_document(svc, text=resume("Cy D", "Accountant", "GST, Tally"))
    page = await search(svc, "python django backend", limit=3)
    assert page.total == 3 and page.candidates[0].name == "Asha R"
    first = page.candidates[0]
    assert first.keyword_rank == 1 and first.meaning_rank == 1 and "Python" in first.profile
    assert page.candidates[-1].keyword_rank is None  # listed on meaning alone


async def test_pages_and_new_people_reach_the_index(svc):
    for name, head, skills in [("Ana P", "Systems engineer", "Rust, Linux kernel, eBPF"),
                               ("Ben Q", "Embedded developer", "Rust, C, RTOS firmware"),
                               ("Cat R", "Blockchain engineer", "Rust, Solana, smart contracts")]:
        await add_document(svc, text=resume(name, head, skills))
    p1 = await search(svc, "rust")
    assert [c.position for c in p1.candidates] == [1, 2] and p1.next_offset == 2
    p2 = await search(svc, "rust", offset=p1.next_offset)
    assert len(p2.candidates) == 1 and p2.next_offset is None
    seen = {c.person_id for c in p1.candidates + p2.candidates}
    assert len(seen) == 3
    await add_document(svc, text=resume("Zed Q", "Engineer", "Elixir"))
    assert (await search(svc, "elixir")).candidates[0].name == "Zed Q"


async def test_empty_pool_and_bad_queries(svc):
    page = await search(svc, "anything")
    assert page.total == 0 and page.next_offset is None
    for q, off in [(" ", 0), ("x", -1)]:
        with pytest.raises(ToolFailure) as e:
            await search(svc, q, off)
        assert e.value.code == "invalid_input"


async def test_a_job_searches_with_its_posting_or_the_agents_query(svc):
    await add_document(svc, text=resume("Asha R", "Backend engineer", "Python, Django"))
    await add_document(svc, text=resume("Cy D", "Accountant", "GST, Tally"))
    j = await post_job(svc, "# Senior accountant\nGST filings and Tally, Mumbai office.")
    assert j.title == "Senior accountant"
    calls = svc.embedder.calls
    page = await find_candidates(svc, j.job_id)
    assert page.candidates[0].name == "Cy D" and svc.embedder.calls == calls  # the posting's own vector
    page = await find_candidates(svc, j.job_id, query="python django")
    assert page.candidates[0].name == "Asha R"
    with pytest.raises(ToolFailure) as e:
        await find_candidates(svc, "f" * 32)
    assert e.value.code == "not_found"


async def test_the_same_person_filed_twice_is_listed_once(svc):
    long = " ".join(f"Built service {n} in Python with Postgres and Redis for payments." for n in range(30))
    a = await add_document(svc, text=resume("Asha R", "Backend engineer", "Python", extra=long))
    b = await add_document(svc, text=resume("Asha Rao", "Senior backend engineer", "Python",
                                            extra=long + " Led a team of four."))
    c = await add_document(svc, text=resume("Ravi K", "Data analyst", "SQL", phone="+91 99999 11111"))
    d = await add_document(svc, text=resume("R Kumar", "Analyst", "Excel", phone="+91-99999-11111"))
    await add_document(svc, text=resume("Bo C", "Frontend engineer", "React"))
    page = await search(svc, "python payments analyst", limit=10)
    assert page.total == 3
    listed = {x.person_id: x.duplicates for x in page.candidates}
    assert (a.person_id in listed) != (b.person_id in listed)  # one of the two resumes, the other under it
    assert listed.get(a.person_id, listed.get(b.person_id)) in ([a.person_id], [b.person_id])
    assert (c.person_id in listed) != (d.person_id in listed)  # same phone, written two ways


def test_duplicate_groups_leave_different_people_apart():
    words = [f"w{n}" for n in range(400)]
    one = " ".join(words)
    g = duplicate_groups([one, one.replace("w200", "x"), " ".join(reversed(words)), "short note", "short note"],
                         [None] * 5, [None] * 5)
    assert g[0] == g[1] and g[0] != g[2]  # near-identical text joins, the same words in another order do not
    assert g[3] != g[4]  # too little text to tell
    assert len(set(duplicate_groups(["x"] * 2, ["Same@x.com", "same@x.com"], [None, None]))) == 1


async def test_blank_names_written_by_the_model_are_stored_as_none(svc):
    r = await add_document(svc, text="Name: null\nEmail: unknown\nHeadline: Engineer\nSkill: Go")
    p = get_person(svc, r.person_id)
    assert p.name is None and p.email is None


# ------------------------------------------------------------------ judging a page


def test_check_score_weights_the_first_must_have_twice_and_counts_years():
    full = ("yes", 5.0, 5.0, "yes", "yes", "yes")
    assert check_score(*full, ["yes", "yes", "yes"]) == 100
    assert check_score(*full, []) == 100
    # the first must-have alone weighs as much as the other two together: 25 * 2/4 off either way
    assert check_score(*full, ["yes", "no", "no"]) == check_score(*full, ["no", "yes", "yes"]) == 87.5
    # years against the years needed: 2 of 4 is half the 15 points; past the need counts as full
    assert check_score("yes", 2.0, 4.0, "yes", "yes", "yes", ["yes"]) == 92.5
    assert check_score("yes", 9.0, 4.0, "yes", "yes", "yes", ["yes"]) == 100
    assert check_score("yes", 1.5, None, "yes", "yes", "yes", ["yes"]) == 92.5  # no need given: 3 years
    assert check_score("no", 9.0, 4.0, "yes", "yes", "yes", ["yes"]) == 55  # no core work: its years count 0


def test_search_place_breaks_ties():
    assert search_score(1, 477) == 100 and search_score(477, 477) == 0 and search_score(1, 1) == 100
    assert score(60, search_score(2, 100)) > score(60, search_score(40, 100))
    assert score(80, 40) == 74.0


async def test_judging_orders_the_page_by_score_then_name(svc):
    people = [("Zoe A", "Core: yes\nRole: yes\nMusts: yes,yes,maybe"),
              ("Amy B", "Core: yes\nRole: yes\nMusts: yes,yes,maybe"),
              ("Max C", "Core: no\nRole: yes\nMusts: yes,yes,yes"),
              ("Kim D", "Core: maybe\nRole: maybe\nMusts: yes"),
              ("Lee E", "Core: yes\nRole: no\nMusts: yes,yes,yes"),
              ("Sam F", "Core: yes\nRole: yes\nMusts: yes,yes,yes")]
    skills = ["Python, Django", "Python, Flask", "Python, Pandas", "Python, FastAPI", "Python, Celery",
              "Python, Postgres"]
    for (name, checks), sk in zip(people, skills):
        await add_document(svc, text=resume(name, f"{name.split()[0]} backend engineer", sk, extra=checks))
    job = await post_job(svc, "Backend engineer, Python and Postgres")
    page = await find_candidates(svc, job.job_id, limit=6)
    judged = await judge_page(svc, job.job_id, page)
    by_name = {c.name: c for c in judged.candidates}
    assert [c.name for c in judged.candidates][0] == "Sam F"
    assert [c.name for c in judged.candidates][-3:] == ["Lee E", "Kim D", "Max C"]
    for c in judged.candidates:  # the score is the checks and the search place, as stored with it
        assert c.search_score == round(search_score(c.position, page.total), 1)
    amy, zoe = by_name["Amy B"], by_name["Zoe A"]
    assert amy.check_score == zoe.check_score  # same answers: the search place decides
    assert (amy.score > zoe.score) == (amy.position < zoe.position)
    kim = by_name["Kim D"]
    assert kim.must_haves == {"Services in production": "yes", "Python": "no", "Postgres": "no"}  # unanswered: no
    assert judged.must_haves == ["Services in production", "Python", "Postgres"]
    assert svc.judge.calls == ["job_core"] + ["checks"] * 6  # one candidate per call
    await judge_page(svc, job.job_id, page)
    assert svc.judge.calls.count("job_core") == 1  # kept with the job


async def test_a_candidate_the_judge_leaves_out_is_asked_again_then_listed_last(svc):
    a = await add_document(svc, text=resume("Ann A", "Backend engineer", "Python", extra="Core: yes\nRole: yes"))
    b = await add_document(svc, text=resume("Bob B", "Backend engineer", "Go, Python", extra="Core: no"))
    job = await post_job(svc, "Backend engineer")
    svc.judge.skip = {a.person_id[:8]}
    judged = await judge_page(svc, job.job_id, await find_candidates(svc, job.job_id))
    assert [c.person_id for c in judged.candidates] == [b.person_id, a.person_id]
    assert judged.candidates[-1].score is None
    assert svc.judge.calls == ["job_core", "checks", "checks", "checks"]  # the missing one asked once more


# ------------------------------------------------------------------ the caller's key, never the server's


def test_settings_have_no_api_key_even_when_the_environment_does(monkeypatch):
    from matchmaker.config import Settings

    monkeypatch.setenv("LLM_API_KEY", "sk-server")
    monkeypatch.setenv("EMBED_API_KEY", "sk-server")
    monkeypatch.setenv("LLM_MODEL", "openai/gpt-5-mini")  # the recruiter server's setting in the shared .env
    s = Settings()
    dumped = s.model_dump()
    assert not any("key" in name for name in dumped)
    assert "sk-server" not in json.dumps(dumped, default=str)
    assert s.llm_model == "openai/gpt-5-nano"


def test_model_calls_without_the_callers_key_are_refused(svc):
    from matchmaker.services import Services

    bare = Services(settings=svc.settings)
    for client in ("llm", "embedder", "judge"):
        with pytest.raises(ToolFailure) as e:
            getattr(bare, client)
        assert e.value.code == "missing_header"


def test_the_callers_key_builds_the_clients(svc):
    from matchmaker.services import Services

    bare = Services(settings=svc.settings)
    with billing.use_api_key("sk-caller-1"):
        llm = bare.llm
        assert llm.client.api_key == "sk-caller-1" and bare.embedder.client.api_key == "sk-caller-1"
        assert bare.llm is llm  # reused within one key
    with billing.use_api_key("sk-caller-2"):
        assert bare.llm is not llm


def test_paid_tools_need_the_key_header_and_stdio_is_refused():
    with pytest.raises(ToolFailure) as e:
        billing.begin("add_document", None)
    assert e.value.code == "missing_header"
    with pytest.raises(ToolFailure) as e:
        billing.begin("find_candidates", {"Tool-Call-Max-Charge": "1"})
    assert e.value.code == "missing_header" and "x-openrouter-api-key" in e.value.message
    tokens = billing.begin("post_job", {"X-OpenRouter-Api-Key": "sk-caller", "Tool-Call-Max-Charge": "$0.50"})
    try:
        assert billing.current_api_key() == "sk-caller" and billing.current_budget().approved_usd == 0.5
    finally:
        billing.end(tokens)
    assert billing.current_api_key() is None
    free = billing.begin("get_person", None)  # no model calls: no key needed
    billing.end(free)


# ------------------------------------------------------------------ Tool-Call-Max-Charge


def test_max_charge_header_is_optional_and_parsed():
    tokens = billing.begin("add_document", {"X-OpenRouter-Api-Key": "sk"})
    assert billing.current_budget() is None  # no header: no cap
    billing.end(tokens)
    tokens = billing.begin("add_document", {"X-OpenRouter-Api-Key": "sk", "Tool-Call-Max-Charge": "$0.25"})
    assert billing.current_budget().approved_usd == 0.25
    billing.end(tokens)
    with pytest.raises(ToolFailure) as e:
        billing.begin("add_document", {"X-OpenRouter-Api-Key": "sk", "Tool-Call-Max-Charge": "lots"})
    assert e.value.code == "invalid_input"


def test_an_estimate_over_the_max_charge_is_refused_before_spending():
    tokens = billing.begin("add_document", {"X-OpenRouter-Api-Key": "sk", "Tool-Call-Max-Charge": "0.0005"})
    try:
        with pytest.raises(ToolFailure) as e:
            billing.preflight("add_document")  # the $0.001 default
    finally:
        billing.end(tokens)
    assert e.value.code == "cost_limit_exceeded" and "Tool-Call-Max-Charge" in e.value.message


async def test_spend_reaching_the_max_charge_stops_further_model_calls(svc):
    tokens = billing.begin("add_document", {"X-OpenRouter-Api-Key": "sk", "Tool-Call-Max-Charge": "0.0001"})
    try:
        with pytest.raises(ToolFailure) as e:  # the fake profile costs $0.0001, so the embedding is not started
            await add_document(svc, text=resume("A B", "Backend engineer", "Rust"))
    finally:
        billing.end(tokens)
    assert e.value.code == "cost_limit_exceeded" and "reached" in e.value.message
    assert not (svc.settings.data_dir / "people").exists()

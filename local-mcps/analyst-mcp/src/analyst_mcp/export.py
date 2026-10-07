"""The xlsx a spec turns into: a copy of the original with the new sheet first and, when the spec has checks, a
Checks sheet last. Every number in the new sheet is a live formula over the original sheets."""

from pathlib import Path

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .expr import XlError
from .spec import FIRST_COL, TITLE_ROW, Built, chart_rows, excel_format

FONT = "Arial"
NAVY, TOTAL, ZEBRA, LINE, GREY = "1F3A5F", "E8EEF5", "F6F8FB", "BFC9D6", "595959"
PALETTE = ["1F3A5F", "C9A227", "2E7D4F", "9FB3C8", "B5523B", "6B5B95", "3E8E9D", "8C8C8C"]

thin = Side(style="thin", color=LINE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def _font(size=10, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def _fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def write(built: Built, original: Path, target: Path) -> Path:
    vba = original.suffix.lower() == ".xlsm"
    wb = openpyxl.load_workbook(original, keep_vba=vba)
    spec = built.spec
    ws = wb.create_sheet(spec.sheet_name, 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2

    ws.cell(TITLE_ROW, FIRST_COL, spec.title).font = _font(14, True, NAVY)
    if spec.subtitle:
        ws.cell(TITLE_ROW + 1, FIRST_COL, spec.subtitle).font = _font(10, color=GREY)

    widths: dict[int, float] = {}
    for lay in built.sections:
        sec = lay.section
        ncols = len(sec.columns)
        title = ws.cell(lay.title_row, FIRST_COL, sec.title)
        title.font = _font(12, True, NAVY)
        for c in range(ncols):
            ws.cell(lay.title_row, FIRST_COL + c).border = Border(bottom=Side(style="medium", color=NAVY))
        for c, head in enumerate(sec.columns):
            cell = ws.cell(lay.header_row, FIRST_COL + c, head)
            cell.font = _font(10, True, "FFFFFF")
            cell.fill = _fill(NAVY)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="left" if c == 0 else "center", vertical="center", wrap_text=True)
            widths[c] = max(widths.get(c, 0), min(len(head), 40))
        for r, (row_spec, excel, values) in enumerate(zip(sec.rows, lay.excel, lay.values, strict=True)):
            row = lay.first_row + r
            for c, (content, value) in enumerate(zip(excel, values, strict=True)):
                cell = ws.cell(row, FIRST_COL + c, content)
                fmt = excel_format(sec.formats[c] if sec.formats and c < len(sec.formats) else None)
                if fmt:
                    cell.number_format = fmt
                cell.border = BOX
                bold = row_spec.style in ("total", "highlight")
                color = "FFFFFF" if row_spec.style == "highlight" else "000000"
                cell.font = _font(10, bold, color)
                if row_spec.style == "total":
                    cell.fill = _fill(TOTAL)
                elif row_spec.style == "highlight":
                    cell.fill = _fill(NAVY)
                elif r % 2:
                    cell.fill = _fill(ZEBRA)
                numeric = isinstance(value, int | float) and not isinstance(value, bool)
                cell.alignment = Alignment(
                    horizontal="right" if c > 0 and (numeric or isinstance(value, XlError)) else "left",
                    vertical="center",
                    wrap_text=c == 0,
                )
                shown = len(str(value)) if value is not None else 0
                widths[c] = max(widths.get(c, 0), min(shown + (4 if numeric else 0), 60 if c == 0 else 24))
        if lay.note_row:
            ws.cell(lay.note_row, FIRST_COL, sec.note).font = _font(9, color=GREY, italic=True)
        if sec.chart:
            ws.add_chart(_chart(ws, lay), lay.chart_anchor)

    for c, w in widths.items():
        letter = get_column_letter(FIRST_COL + c)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, 12 if c else 18, w + 2)

    if built.checks:
        _checks_sheet(wb, built)

    ws.sheet_properties.tabColor = NAVY
    wb.active = 0
    for sheet in wb.worksheets:
        sheet.sheet_view.tabSelected = sheet is ws
    # openpyxl writes formulas without results; Excel, Google Sheets and Numbers compute them on open.
    wb.calculation.fullCalcOnLoad = True
    target = target.with_suffix(".xlsm" if vba else ".xlsx")
    tmp = target.with_name(target.name + ".tmp")
    wb.save(tmp)
    tmp.replace(target)
    return target


def _chart(ws, lay):
    sec, spec = lay.section, lay.section.chart
    rows = chart_rows(sec)
    if spec.type == "pie":
        chart = PieChart()
    elif spec.type == "line":
        chart = LineChart()
    else:
        chart = BarChart()
        chart.type = "bar" if spec.type == "bar" else "col"
        if spec.type == "bar":
            chart.x_axis.scaling.orientation = "maxMin"
    chart.title = spec.title or sec.title
    chart.width = 17 if spec.type != "pie" else 13
    chart.height = 8.5
    first, last = lay.first_row + rows[0] - 1, lay.first_row + rows[-1] - 1
    for col in spec.value_columns:
        chart.add_data(Reference(ws, min_col=FIRST_COL + col - 1, min_row=first, max_row=last))
    chart.set_categories(Reference(ws, min_col=FIRST_COL + spec.label_column - 1, min_row=first, max_row=last))
    from openpyxl.chart.series import SeriesLabel

    for k, col in enumerate(spec.value_columns):
        chart.series[k].tx = SeriesLabel(v=sec.columns[col - 1])
        if spec.type != "pie":
            chart.series[k].graphicalProperties.solidFill = PALETTE[k % len(PALETTE)]
            chart.series[k].graphicalProperties.line.solidFill = PALETTE[k % len(PALETTE)]
    labels = DataLabelList()
    labels.showVal = spec.type != "pie"
    labels.showPercent = spec.type == "pie"
    labels.showSerName = labels.showCatName = labels.showLegendKey = False
    chart.dataLabels = labels
    if spec.type == "pie":
        from openpyxl.chart.marker import DataPoint

        for i in range(len(rows)):
            point = DataPoint(idx=i)
            point.graphicalProperties.solidFill = PALETTE[i % len(PALETTE)]
            chart.series[0].dPt.append(point)
    else:
        chart.y_axis.delete = False
        chart.x_axis.delete = False
        if len(spec.value_columns) == 1:
            chart.legend = None
    return chart


def _checks_sheet(wb, built: Built) -> None:
    name = "Checks"
    taken = {s.lower() for s in wb.sheetnames}
    n = 2
    while name.lower() in taken:
        name = f"Checks {n}"
        n += 1
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (52, 18, 18, 14, 10), strict=True):
        ws.column_dimensions[col].width = w
    ws["A1"] = f"Checks for the {built.spec.sheet_name} sheet"
    ws["A1"].font = _font(13, True, NAVY)
    for c, head in enumerate(("Check", "Left", "Right", "Difference", "Result"), 1):
        cell = ws.cell(3, c, head)
        cell.font = _font(10, True, "FFFFFF")
        cell.fill = _fill(NAVY)
        cell.border = BOX
    for i, res in enumerate(built.checks):
        r = 4 + i
        row = [res.check.label, res.left_excel, res.right_excel, f"=IFERROR(B{r}-C{r},\"\")",
               f'=IF(IFERROR(ABS(B{r}-C{r})<={res.check.tolerance},EXACT(B{r},C{r})),"OK","CHECK")']  # fmt: skip
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = _font(10, bold=c == 5, color="2E7D4F" if c == 5 else "000000")
            cell.border = BOX
            if c in (2, 3, 4):
                cell.number_format = "#,##0.00"
    last = 3 + len(built.checks)
    ws.cell(last + 2, 1, "All checks").font = _font(11, True, NAVY)
    ws.cell(last + 2, 2, f'=IF(COUNTIF(E4:E{last},"CHECK")=0,"ALL OK","SEE ABOVE")').font = _font(11, True, "2E7D4F")
    ws.sheet_properties.tabColor = "9FB3C8"

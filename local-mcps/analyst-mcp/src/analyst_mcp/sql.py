"""Read-only SQL over a project's files with DuckDB.

Tables:
  cells                  every non-empty cell of every file: file, sheet, cell, row, col, value, number, formula,
                         label, header, unit
  "<sheet name>"         a sheet read as a table when it has a header row (text across the top, data below). A sheet
                         name several files share is all of them stacked, with a `file` column (May, June, July)
  "<file>.<sheet name>"  one file's sheet, when the project has several files

The database is in memory and built from the files; external access (files, network) is switched off before any
query runs, and only reading statements are accepted.
"""

import re
import threading

import duckdb
import pandas as pd
from openpyxl.utils import get_column_letter

from .errors import ToolFailure
from .workbook import Workbook, show

READING = re.compile(r"^\s*(select|with|from|describe|show|summarize|pivot|unpivot|values|table)\b", re.IGNORECASE)

_conns: dict[str, tuple[duckdb.DuckDBPyConnection, list[str]]] = {}
_lock = threading.Lock()


def _cells_frame(wb: Workbook) -> pd.DataFrame:
    rows = []
    for cell in wb.cells.values():
        fact = wb.fact(cell)
        number = cell.value if isinstance(cell.value, int | float) and not isinstance(cell.value, bool) else None
        rows.append(
            {
                "sheet": cell.sheet,
                "cell": cell.coord,
                "row": cell.row,
                "col": get_column_letter(cell.col),
                "value": show(cell.value),
                "number": float(number) if number is not None else None,
                "formula": cell.formula,
                "label": fact.label if number is not None else None,
                "header": fact.header if number is not None else None,
                "unit": fact.unit if number is not None else None,
            }
        )
    frame = pd.DataFrame(rows, columns=["sheet", "cell", "row", "col", "value", "number", "formula", "label",
                                        "header", "unit"])  # fmt: skip
    return frame.astype({"number": "float64"})


def _sheet_frame(wb: Workbook, sheet: str) -> pd.DataFrame | None:
    """The sheet as a table, when one row of text heads at least two columns of data below it."""
    cells = wb.sheet_cells(sheet)
    if not cells:
        return None
    by_row: dict[int, dict[int, object]] = {}
    for c in cells:
        by_row.setdefault(c.row, {})[c.col] = c.value
    rows = sorted(by_row)
    for i, r in enumerate(rows[:10]):
        head = by_row[r]
        if len(head) >= 2 and all(isinstance(v, str) for v in head.values()) and i + 1 < len(rows):
            below = rows[i + 1 :]
            if len(below) < 2:
                return None
            cols = sorted(head)
            names, seen = [], set()
            for col in cols:
                name = re.sub(r"\s+", " ", str(head[col])).strip() or get_column_letter(col)
                base, n = name, 2
                while name.lower() in seen:
                    name = f"{base} {n}"
                    n += 1
                seen.add(name.lower())
                names.append(name)
            data = [[by_row[rr].get(col) for col in cols] for rr in below]
            frame = pd.DataFrame(data, columns=names)
            for name in names:
                converted = pd.to_numeric(frame[name], errors="coerce")
                if converted.notna().sum() >= max(1, int(frame[name].notna().sum() * 0.8)):
                    frame[name] = converted
                else:
                    frame[name] = frame[name].map(lambda v: None if v is None else show(v))
            return frame
    return None


def _table(con: duckdb.DuckDBPyConnection, name: str, frame: pd.DataFrame) -> None:
    con.register("frame", frame)
    con.execute(f'CREATE TABLE "{name.replace(chr(34), chr(34) * 2)}" AS SELECT * FROM frame')
    con.unregister("frame")


def connection(key: str, books: list[tuple[str, Workbook]]) -> tuple[duckdb.DuckDBPyConnection, list[str]]:
    """The project's database: `books` are (file name, loaded workbook); `key` names this set of versions."""
    with _lock:
        if key in _conns:
            return _conns[key]
        con = duckdb.connect(":memory:")
        tables = ["cells"]
        frames = []
        for name, wb in books:
            frame = _cells_frame(wb)
            frame.insert(0, "file", name)
            frames.append(frame)
        _table(con, "cells", pd.concat(frames, ignore_index=True))
        by_sheet: dict[str, list[tuple[str, pd.DataFrame]]] = {}
        for name, wb in books:
            for sheet in wb.sheets:
                sf = _sheet_frame(wb, sheet)
                if sf is not None and sheet.lower() != "cells":
                    by_sheet.setdefault(sheet, []).append((name, sf))
        several = len(books) > 1
        for sheet, parts in by_sheet.items():
            if several:
                for name, sf in parts:
                    _table(con, f"{name}.{sheet}", sf)
                    tables.append(f"{name}.{sheet}")
                stacked = pd.concat([sf.assign(file=name)[["file", *sf.columns]] for name, sf in parts],
                                    ignore_index=True)  # fmt: skip
                _table(con, sheet, stacked)
            else:
                _table(con, sheet, parts[0][1])
            tables.append(sheet)
        con.execute("SET enable_external_access = false")
        con.execute("SET lock_configuration = true")
        _conns[key] = (con, tables)
        return con, tables


def forget(project_key: str, keep: str = "") -> None:
    """Close the project's databases (one per set of file versions), all but `keep`."""
    with _lock:
        gone = [_conns.pop(k) for k in list(_conns) if k.split(":")[0] == project_key and k != keep]
    for con, _ in gone:
        con.close()


def describe(key: str, books: list[tuple[str, Workbook]]) -> str:
    con, tables = connection(key, books)
    lines = []
    for t in tables:
        cols = con.cursor().execute(f'DESCRIBE "{t}"').fetchall()
        count = con.cursor().execute(f'SELECT count(*) FROM "{t}"').fetchone()[0]
        lines.append(f'- "{t}" ({count} rows): ' + ", ".join(f'"{c[0]}" {c[1]}' for c in cols))
    return "\n".join(lines)


def run(key: str, books: list[tuple[str, Workbook]], sql: str, limit: int) -> tuple[list[str], list[tuple], bool]:
    text = sql.strip().rstrip(";").strip()
    if not text or ";" in text:
        raise ToolFailure("bad_sql", "send exactly one statement")
    if not READING.match(text):
        raise ToolFailure("read_only", "only reading statements run here (SELECT, WITH, DESCRIBE, SUMMARIZE, …)")
    con, _ = connection(key, books)
    cur = con.cursor()
    try:
        cur.execute(text)
    except duckdb.Error as e:
        raise ToolFailure("sql_error", str(e).split("\n")[0][:400]) from e
    columns = [d[0] for d in cur.description or []]
    rows = cur.fetchmany(limit + 1)
    return columns, rows[:limit], len(rows) > limit

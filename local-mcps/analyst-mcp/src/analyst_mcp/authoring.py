"""Workbooks the agent writes: new ones from rows, and edited copies of open ones.

A cell value starting with "=" is a formula. Values are written as given, so numbers stay numbers. The file an edit
starts from is never written: the result is new bytes, opened as a workbook of its own.
"""

import io
import re
from copy import copy
from pathlib import Path
from typing import Literal

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, Field

from .errors import ToolFailure
from .expr import quote_sheet

CellValue = str | int | float | bool | None
Rows = list[list[CellValue]]

BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
HEX = re.compile(r"^#?([0-9A-Fa-f]{6})$")


class NewSheet(BaseModel):
    name: str = Field(description="Tab name, up to 31 characters")
    rows: Rows = Field(description="Rows of cells, the first row the column headings. '=...' is a formula")


class Edit(BaseModel):
    """One change. Which fields count depends on op."""

    op: Literal["set", "append_rows", "add_sheet", "rename_sheet", "delete_sheet", "format"] = Field(
        description=(
            "set: write `cells` into `sheet`. append_rows: add `rows` under the last used row of `sheet`. "
            "add_sheet: new sheet `sheet`, optionally with `rows`. rename_sheet: `sheet` becomes `to` (formulas "
            "follow). delete_sheet: remove `sheet`. format: style `range` of `sheet` (bold, number_format, fill, "
            "width)"
        )
    )
    sheet: str = Field(description="The sheet the change applies to")
    cells: dict[str, CellValue] | None = Field(
        None, description='set: address to value, e.g. {"B3": 1200, "C3": "=B3*1.18", "A9": null}'
    )
    rows: Rows | None = Field(None, description="append_rows, add_sheet: rows of cells")
    to: str | None = Field(None, description="rename_sheet: the new name")
    range: str | None = Field(None, description="format: cells to style, e.g. A1:F1 or B:B")
    bold: bool | None = Field(None, description="format: bold on or off")
    number_format: str | None = Field(None, description="format: Excel number format, e.g. #,##0.00 or 0%")
    fill: str | None = Field(None, description="format: background colour as hex, e.g. FFF2CC")
    width: float | None = Field(None, gt=0, le=255, description="format: column width for the range's columns")


def _sheet_name(name: str, taken: list[str]) -> str:
    cleaned = BAD_SHEET_CHARS.sub(" ", name).strip()[:31]
    if not cleaned:
        raise ToolFailure("bad_sheet_name", f"{name!r} is not a usable sheet name")
    if cleaned.lower() in (t.lower() for t in taken):
        raise ToolFailure("sheet_exists", f"there is already a sheet named {cleaned!r}")
    return cleaned


def _write_rows(ws: Worksheet, rows: Rows, first_row: int) -> None:
    for r, row in enumerate(rows):
        if not isinstance(row, list):
            raise ToolFailure("bad_rows", f"row {r + 1} is not a list of cells")
        for c, value in enumerate(row):
            if value is not None and value != "":
                ws.cell(first_row + r, c + 1, value)


def _fit_columns(ws: Worksheet) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or (isinstance(cell.value, str) and cell.value.startswith("=")):
                continue
            longest = max(len(line) for line in str(cell.value).split("\n"))
            widths[cell.column] = max(widths.get(cell.column, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width, 6), 60) + 2


def create(sheets: list[NewSheet]) -> bytes:
    """A new workbook: each sheet's first row bold and frozen, columns sized to their text."""
    if not sheets:
        raise ToolFailure("no_sheets", "a workbook needs at least one sheet")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet in sheets:
        ws = wb.create_sheet(_sheet_name(sheet.name, wb.sheetnames))
        _write_rows(ws, sheet.rows, 1)
        if sheet.rows:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
        _fit_columns(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _find(wb: openpyxl.Workbook, name: str) -> Worksheet:
    for title in wb.sheetnames:
        if title.lower() == name.strip().lower():
            return wb[title]
    raise ToolFailure("sheet_not_found", f"no sheet {name!r}; sheets: " + ", ".join(wb.sheetnames))


def _reference(name: str) -> re.Pattern:
    """A formula's reference to sheet `name`, quoted or not."""
    quoted = "'" + re.escape(name.replace("'", "''")) + "'!"
    bare = r"(?<![A-Za-z0-9_.'])" + re.escape(name) + "!"
    return re.compile(f"{quoted}|{bare}", re.IGNORECASE)


def _formulas(wb: openpyxl.Workbook):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    yield ws, cell


def _set(ws: Worksheet, cells: dict[str, CellValue]) -> list[tuple[str, str]]:
    done = []
    for address, value in cells.items():
        coord = address.replace("$", "").strip().upper()
        try:
            coordinate_from_string(coord)
        except ValueError as e:
            raise ToolFailure("bad_cell", f"{address!r} is not a cell address like B3") from e
        ws[coord] = None if value == "" else value
        done.append((ws.title, coord))
    return done


def _format(ws: Worksheet, edit: Edit) -> None:
    if not edit.range:
        raise ToolFailure("bad_edit", "format needs a range, e.g. A1:F1")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(edit.range.replace("$", "").upper())
    except ValueError as e:
        raise ToolFailure("bad_range", f"{edit.range!r} is not a range like A1:F40") from e
    min_col, min_row = min_col or 1, min_row or 1
    max_col, max_row = max_col or ws.max_column, max_row or ws.max_row
    fill = None
    if edit.fill:
        match = HEX.match(edit.fill.strip())
        if not match:
            raise ToolFailure("bad_colour", f"{edit.fill!r} is not a hex colour like FFF2CC")
        fill = PatternFill("solid", start_color=match.group(1).upper(), end_color=match.group(1).upper())
    if edit.bold is not None or edit.number_format or fill:
        for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            for cell in row:
                if edit.bold is not None:
                    font = copy(cell.font)
                    font.bold = edit.bold
                    cell.font = font
                if edit.number_format:
                    cell.number_format = edit.number_format
                if fill:
                    cell.fill = fill
    if edit.width:
        for col in range(min_col, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = edit.width


def apply(original: Path, edits: list[Edit], macros: bool) -> tuple[bytes, list[tuple[str, str]]]:
    """The workbook with `edits` applied in order, and the cells `set` wrote. Any edit that cannot be made refuses
    the lot, so the user never gets half a change."""
    if not edits:
        raise ToolFailure("no_edits", "pass at least one edit")
    wb = openpyxl.load_workbook(original, keep_vba=macros)
    written: list[tuple[str, str]] = []
    for n, edit in enumerate(edits, 1):
        try:
            if edit.op == "add_sheet":
                ws = wb.create_sheet(_sheet_name(edit.sheet, wb.sheetnames))
                _write_rows(ws, edit.rows or [], 1)
                _fit_columns(ws)
                continue
            ws = _find(wb, edit.sheet)
            if edit.op == "set":
                if not edit.cells:
                    raise ToolFailure("bad_edit", "set needs cells")
                written += _set(ws, edit.cells)
            elif edit.op == "append_rows":
                if not edit.rows:
                    raise ToolFailure("bad_edit", "append_rows needs rows")
                start = ws.max_row + 1 if ws.max_row > 1 or ws["A1"].value is not None else 1
                _write_rows(ws, edit.rows, start)
            elif edit.op == "rename_sheet":
                if not edit.to:
                    raise ToolFailure("bad_edit", "rename_sheet needs to")
                new = _sheet_name(edit.to, [s for s in wb.sheetnames if s != ws.title])
                pattern = _reference(ws.title)
                for _, cell in _formulas(wb):
                    cell.value = pattern.sub(lambda _m: quote_sheet(new) + "!", cell.value)
                ws.title = new
            elif edit.op == "delete_sheet":
                if len(wb.sheetnames) == 1:
                    raise ToolFailure("last_sheet", "a workbook must keep at least one sheet")
                pattern = _reference(ws.title)
                users = [
                    f"{other.title}!{cell.coordinate}"
                    for other, cell in _formulas(wb)
                    if other is not ws and pattern.search(cell.value)
                ]
                if users:
                    raise ToolFailure(
                        "sheet_in_use",
                        f"formulas still use {ws.title!r}: " + ", ".join(users[:10]) + "; change them first",
                    )
                wb.remove(ws)
            elif edit.op == "format":
                _format(ws, edit)
        except ToolFailure as e:
            raise ToolFailure(e.code, f"edit {n} ({edit.op}): {e.message}") from e
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), written

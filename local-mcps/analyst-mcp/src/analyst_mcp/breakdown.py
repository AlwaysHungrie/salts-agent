"""What a total is made of: the cells that add up to it.

A total like `=C92+C35` with `C92 = SUM(C38:C91)-C46-C57` is opened up, level by level, into the items that really
sum to it, so the agent can group them into heads without guessing. A cell is opened up only when its own formula
is a plain sum (cells added and subtracted, SUM of ranges) and opening it keeps every item counted once; a cell
whose formula multiplies or adds constants (`=C35*7%`) is an item in its own right.
"""

from dataclasses import dataclass

from openpyxl.utils import range_boundaries

from . import expr
from .errors import ToolFailure
from .workbook import Workbook

Key = tuple[str, str]


@dataclass
class Item:
    sheet: str
    coord: str
    sign: int
    value: float
    label: str
    formula: str | None


def qualify(formula: str, sheet: str) -> str:
    """The workbook's own formula with every bare reference given its sheet, so expr can read it."""
    body = formula[1:]
    out, pos = [], 0
    while pos < len(body):
        m = expr.TOKEN.match(body, pos)
        if not m:
            raise ToolFailure("not_a_sum", f"cannot read {formula!r}")
        out.append(f"{expr.quote_sheet(sheet)}!{m.group(0)}" if m.lastgroup == "bare" else m.group(0))
        pos = m.end()
    return "=" + "".join(out)


def _linear(node, wb: Workbook) -> dict[Key, int] | None:
    """The formula as cells with +1/-1 weights, or None when it is not a plain sum of cells."""
    kind = node[0]
    if kind == "ref":
        text = node[1]
        sheet, _, rng = text.rpartition("!")
        sheet = wb.sheet(sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet)
        rng = rng.replace("$", "").upper()
        if ":" not in rng:
            return {(sheet, rng): 1}
        c1, r1, c2, r2 = range_boundaries(rng)
        return {(sheet, expr.address(c, r)): 1 for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)}
    if kind == "fn" and node[1] == "SUM":
        out: dict[Key, int] = {}
        for arg in node[2]:
            part = _linear(arg, wb)
            if part is None:
                return None
            _add(out, part, 1)
        return out
    if kind == "neg":
        part = _linear(node[1], wb)
        return None if part is None else {k: -v for k, v in part.items()}
    if kind == "bin" and node[1] in ("+", "-"):
        left, right = _linear(node[2], wb), _linear(node[3], wb)
        if left is None or right is None:
            return None
        out = dict(left)
        _add(out, right, 1 if node[1] == "+" else -1)
        return out
    return None


def _add(into: dict[Key, int], part: dict[Key, int], k: int) -> None:
    for key, w in part.items():
        into[key] = into.get(key, 0) + k * w


def _numeric(wb: Workbook, key: Key) -> float | None:
    v = wb.value(*key)
    return float(v) if isinstance(v, int | float) and not isinstance(v, bool) else None


def _of(wb: Workbook, key: Key) -> dict[Key, int] | None:
    cell = wb.cell(*key)
    if cell is None or not cell.formula:
        return None
    try:
        return _linear(expr.parse(qualify(cell.formula, key[0])), wb)
    except ToolFailure:
        return None


def components(wb: Workbook, sheet: str, coord: str) -> tuple[list[Item], float]:
    """The items that add up to the cell, and the cell's value."""
    key = (sheet, coord)
    total = _numeric(wb, key)
    if total is None:
        raise ToolFailure("not_a_number", f"{sheet}!{coord} holds no number")
    terms = _of(wb, key)
    if terms is None:
        cell = wb.cell(*key)
        raise ToolFailure(
            "not_a_sum",
            f"{sheet}!{coord} is not a sum of other cells (formula: {cell.formula if cell else 'none'}); "
            "pick the total cell itself",
        )
    opened = {key}
    changed = True
    while changed:
        changed = False
        for k in list(terms):
            w = terms.get(k, 0)
            if w == 0 or k in opened:
                continue
            parts = _of(wb, k)
            opened.add(k)
            # A plain link (=Other!B32) stays one item: its label here says what it is.
            if not parts or len(parts) == 1:
                continue
            trial = dict(terms)
            trial.pop(k)
            _add(trial, parts, w)
            # Opening it must keep every item counted once (no item twice, no item that is not a number).
            live = {kk: ww for kk, ww in trial.items() if ww != 0 and _numeric(wb, kk) not in (None, 0.0)}
            if all(abs(ww) == 1 for ww in live.values()):
                terms = trial
                changed = True
    items = []
    for (s, c), w in terms.items():
        v = _numeric(wb, (s, c))
        if w == 0 or v is None or v == 0:
            continue
        cell = wb.cell(s, c)
        items.append(Item(s, c, w, v, wb.fact(cell).label, cell.formula))
    items.sort(key=lambda it: (wb.sheets.index(it.sheet), wb.cell(it.sheet, it.coord).row))
    return items, total


def assign(wb: Workbook, items: list[Item], groups: dict[str, list[str]], default_sheet: str) -> dict[str, list[Item]]:
    """Each head's items. Every item must land in exactly one head, and heads may name only items of the total."""
    by_key = {(it.sheet, it.coord): it for it in items}
    out: dict[str, list[Item]] = {}
    seen: dict[Key, str] = {}
    problems = []
    for name, refs in groups.items():
        out[name] = []
        for ref in refs:
            for key in _refs(wb, ref, default_sheet):
                if key not in by_key:
                    if _numeric(wb, key) not in (None, 0.0):
                        problems.append(f"{key[0]}!{key[1]} (in {name!r}) is not one of the items")
                    continue
                if key in seen:
                    problems.append(f"{key[0]}!{key[1]} is in both {seen[key]!r} and {name!r}")
                    continue
                seen[key] = name
                out[name].append(by_key[key])
    missing = [f"{it.coord} {it.label[:40]}" for it in items if (it.sheet, it.coord) not in seen]
    if missing:
        problems.append("not in any head: " + "; ".join(missing))
    empty = [name for name, its in out.items() if not its]
    if empty:
        problems.append("heads with no items: " + ", ".join(repr(n) for n in empty))
    if problems:
        raise ToolFailure("bad_groups", " | ".join(problems))
    return out


def _refs(wb: Workbook, ref: str, default_sheet: str) -> list[Key]:
    sheet, _, rng = ref.strip().rpartition("!")
    sheet = wb.sheet(sheet[1:-1].replace("''", "'") if sheet.startswith("'") else sheet) if sheet else default_sheet
    rng = rng.replace("$", "").upper()
    try:
        c1, r1, c2, r2 = range_boundaries(rng if ":" in rng else f"{rng}:{rng}")
    except ValueError as e:
        raise ToolFailure("bad_reference", f"{ref!r} is not a cell like C33 or a range like C74:C77") from e
    return [(sheet, expr.address(c, r)) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]


def formula(items: list[Item]) -> str:
    """=SUM(a,b,…)-c-… over the head's items, sheet-qualified as specs need."""
    def ref(it: Item) -> str:
        return f"{expr.quote_sheet(it.sheet)}!{it.coord}"

    plus = [ref(it) for it in items if it.sign > 0]
    minus = [ref(it) for it in items if it.sign < 0]
    head = ref(next(it for it in items if it.sign > 0)) if len(plus) == 1 else f"SUM({','.join(plus)})" if plus else "0"
    return "=" + head + "".join(f"-{m}" for m in minus)

import base64
import json
from pathlib import Path

import openpyxl
import pytest

from analyst_mcp import config, server


@pytest.fixture
def settings(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("MCP_AUTH_TOKEN", "test-token")
    config.get_settings.cache_clear()
    server._books.clear()
    s = config.get_settings()
    s.inbox().mkdir(parents=True, exist_ok=True)
    yield s
    config.get_settings.cache_clear()


def make_model(path: Path) -> Path:
    """A small feasibility-style model. Saved by openpyxl, so it carries formulas but no cached results, which makes
    the server recalculate it."""
    wb = openpyxl.Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    inputs["A1"], inputs["B1"], inputs["C1"] = "Plot Area", 1000, "Sq.mt."
    inputs["A2"], inputs["B2"], inputs["C2"] = "Sale rate", 30000, "Rs."
    inputs["A3"], inputs["B3"] = "Cost per sq ft", 4000
    model = wb.create_sheet("Model Sheet")
    model["A1"], model["B1"] = "Saleable area", "=Inputs!B1*2*10.764*0.9"
    model["A2"], model["B2"] = "Revenue", "=B1*Inputs!B2"
    model["A3"], model["B3"] = "Construction", "=B1*Inputs!B3"
    model["A4"], model["B4"] = "Approvals", 5000000
    model["A5"], model["B5"] = "Total cost", "=SUM(B3:B4)"
    model["A6"], model["B6"] = "Net profit", "=B2-B5"
    tenants = wb.create_sheet("Tenants")
    tenants.append(["Tenant", "Existing", "New"])
    for i, area in enumerate([25.5, 27.0, 26.1, 28.2, 24.9, 26.6], 1):
        tenants.append([f"T{i}", area, f"=B{i + 1}*1.1"])
    tenants["B8"] = "=SUM(B2:B7)"
    wb.save(path)
    return path


@pytest.fixture
def model_file(settings) -> Path:
    return make_model(settings.inbox() / "model.xlsx")


@pytest.fixture
async def proj(model_file) -> str:
    res = await server.open_file(inbox_file=model_file.name)
    assert not res.is_error, res.content[0].text
    return "model"


async def build_draft(proj: str, spec: dict) -> list:
    """Adds a whole spec to an empty draft a section at a time, the spec's checks with the last section. Stops at the
    first error, which is then the last result."""
    from analyst_mcp import store

    files = store.get(config.get_settings(), proj)
    (files.dir / "draft.json").unlink(missing_ok=True)
    out = []
    last = len(spec["sections"])
    for i, section in enumerate(spec["sections"], 1):
        res = await server.add_section(
            proj, section, checks=spec.get("checks") if i == last else None, title=spec.get("title"),
            sheet_name=spec.get("sheet_name"),
        )  # fmt: skip
        out.append(res)
        if res.is_error:
            break
    return out


def text(res) -> str:
    return res.content[0].text


def error(res) -> dict:
    assert res.is_error, text(res)
    return json.loads(text(res))["error"]


def blobs(res, kind: str) -> list[bytes]:
    if kind == "image":
        return [base64.b64decode(c.data) for c in res.content if c.type == "image"]
    return [base64.b64decode(c.resource.blob) for c in res.content if c.type == "resource"]


SAMPLE_SPEC = {
    "sheet_name": "Dashboard",
    "title": "Project summary",
    "sections": [
        {
            "title": "1. Area",
            "columns": ["Item", "Sq. mt.", "Sq. ft."],
            "formats": [None, "num2", "int"],
            "rows": [["Plot area", "=Inputs!B1", "=[r1c2]*10.764"]],
        },
        {
            "title": "2. Profit",
            "columns": ["Item", "Amount", "₹ Crore"],
            "formats": [None, "inr", "num2"],
            "rows": [
                ["Revenue", "='Model Sheet'!B2", "=[r1c2]/10^7"],
                ["Costs", "='Model Sheet'!B5", "=[r2c2]/10^7"],
                {"cells": ["Net profit", "=[r1c2]-[r2c2]", "=[r3c2]/10^7"], "style": "total"},
            ],
            "chart": {"type": "column", "value_columns": [3], "rows": [1, 2, 3]},
        },
    ],
    "checks": [
        {"label": "Profit matches the model", "left": "=[s2r3c2]", "right": "='Model Sheet'!B6", "tolerance": 1}
    ],
}

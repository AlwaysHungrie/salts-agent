import copy
import io
import json

import openpyxl
import pytest

from analyst_mcp import server

from .conftest import SAMPLE_SPEC, blobs, build_draft, error, text

PROFIT = 1000 * 2 * 10.764 * 0.9 * 30000 - (1000 * 2 * 10.764 * 0.9 * 4000 + 5000000)


async def test_open_recalculates_and_maps(proj):
    res = await server.open_file(inbox_file="model.xlsx")
    body = text(res)
    assert 'project "model"' in body
    assert "recalculated" in body
    assert "| B6 | Net profit |" in body
    # A table column is summarised once, its SUM row named as the total.
    assert 'Column "Existing" B2:B7: 6 numbers' in body
    assert "total row B8" in body
    assert '"Tenants"' in body  # also a SQL table


async def test_open_same_file_twice_gives_same_project(proj):
    again = await server.open_file(inbox_file="model.xlsx")
    assert '"model" is already in project "model"' in text(again)


async def test_inbox_name_cannot_escape(settings, proj):
    res = await server.open_file(inbox_file="../projects/x.xlsx")
    assert error(res)["code"] == "inbox_file_not_found"


async def test_find_and_read(proj):
    found = text(await server.find(proj, "net PROFIT"))
    assert "Model Sheet!A6" in found and f"{PROFIT:,.0f}"[:6] in found
    read = text(await server.read_sheet(proj, "model sheet", "A1:B2"))
    assert "B2:" in read and "[=B1*Inputs!B2]" in read


async def test_query_is_read_only_and_charts(proj):
    res = await server.query(proj, 'SELECT Tenant, Existing FROM "Tenants" ORDER BY Existing DESC', chart="bar")
    assert "| T4 | 28.2 |" in text(res)
    assert len(blobs(res, "image")) == 1
    assert error(await server.query(proj, "DROP TABLE cells"))["code"] == "read_only"
    assert error(await server.query(proj, "SELECT 1; SELECT 2"))["code"] == "bad_sql"
    leak = await server.query(proj, "SELECT * FROM read_csv('/etc/hosts')")
    assert error(leak)["code"] == "sql_error"


async def test_query_file(proj):
    res = await server.query(proj, "SELECT sheet, cell, number FROM cells WHERE number IS NOT NULL", file_format="xlsx")
    (data,) = blobs(res, "resource")
    rows = list(openpyxl.load_workbook(io.BytesIO(data)).active.values)
    assert rows[0] == ("sheet", "cell", "number") and len(rows) > 10


async def test_what_if(proj):
    res = text(await server.what_if(proj, {"Inputs!B2": 33000}, ["'Model Sheet'!B6"]))
    expected = PROFIT + 1000 * 2 * 10.764 * 0.9 * 3000
    assert f"{expected:,.0f}"[:7] in res
    assert "+" in res  # relative change shown


async def test_preview_numbers_images_and_checks(proj):
    first, second = await build_draft(proj, SAMPLE_SPEC)
    assert "| Plot area | 1,000.00 | 10,764 |" in text(first)
    assert "- OK: Profit matches the model" in text(second)
    assert f"₹ {PROFIT / 1e7:,.2f}"[:6] not in text(second)  # amount column uses Indian grouping, not crore
    assert len(blobs(first, "image")) == 1 and len(blobs(second, "image")) == 2  # each section once, + its checks


async def test_preview_reports_failing_check(proj):
    spec = copy.deepcopy(SAMPLE_SPEC)
    spec["checks"][0]["right"] = "='Model Sheet'!B5"
    assert "CHECK: Profit matches the model" in text((await build_draft(proj, spec))[-1])
    res = await server.export_sheet(proj)
    assert error(res)["code"] == "checks_failed"


@pytest.mark.parametrize(
    "change, code",
    [
        (lambda s: s["sections"][0]["rows"][0].__setitem__(1, "=B1"), "bad_formula"),
        (lambda s: s.__setitem__("sheet_name", "Inputs"), "sheet_exists"),
        (lambda s: s["sections"][0]["rows"][0].__setitem__(2, "=[r1c3]"), "circular_reference"),
        (lambda s: s["sections"][0]["rows"].append(["a", 1, 2, 3]), "invalid_spec"),
        (lambda s: s["sections"][0]["rows"][0].__setitem__(1, "=Nope!A1"), "sheet_not_found"),
    ],
)
async def test_spec_errors_are_specific(proj, change, code):
    spec = copy.deepcopy(SAMPLE_SPEC)
    change(spec)
    assert error((await build_draft(proj, spec))[-1])["code"] == code


async def test_add_section_advertises_every_field_of_a_section(proj):
    # A model that fills arguments from the schema sent {} while `section` was a bare object.
    (tool,) = [t for t in await server.mcp.list_tools() if t.name == "add_section"]
    props = tool.input_schema["properties"]
    assert "$ref" not in json.dumps(tool.input_schema)
    section = props["section"]["anyOf"][0]
    assert set(section["required"]) == {"title", "columns", "rows"}
    assert {"formats", "chart", "note"} <= set(section["properties"])
    assert props["section"]["description"].startswith("The section")
    check = props["checks"]["anyOf"][0]["items"]
    assert set(check["required"]) == {"label", "left", "right"}
    assert error(await server.add_section(proj, section={}))["code"] == "invalid_spec"


async def test_refused_call_is_logged_with_its_arguments(proj, caplog):
    caplog.set_level("INFO", logger="analyst_mcp")
    res = await server.add_section(proj, section={"title": "Bad", "columns": ["A"], "rows": "not a list"})
    assert error(res)["code"] == "invalid_spec"
    (line,) = [r.getMessage() for r in caplog.records if "add_section failed" in r.getMessage()]
    assert "invalid_spec" in line and "rows" in line
    assert '"not a list"' in line and proj in line


async def test_clear_draft_starts_over_on_the_same_file(proj):
    await build_draft(proj, SAMPLE_SPEC)
    assert not (await server.clear_draft(proj)).is_error
    assert error(await server.export_sheet(proj))["code"] == "no_draft"
    assert "Area" not in text(await server.list_projects())
    res = await server.add_section(proj, SAMPLE_SPEC["sections"][0])
    assert "Draft sections now: " + SAMPLE_SPEC["sections"][0]["title"] + "." in text(res)


async def test_sections_sent_at_once_all_land_in_a_readable_draft(proj):
    """A model sends several add_section calls in one step; they run in threads at the same time. Each must see the
    others' sections, and the draft must stay readable (overlapping writes once left one no tool could open)."""
    import asyncio

    from analyst_mcp import store

    section = SAMPLE_SPEC["sections"][0]
    results = await asyncio.gather(
        *(server.add_section(proj, {**section, "title": f"Part {i}"}) for i in range(8))
    )
    assert not any(r.is_error for r in results), [text(r) for r in results if r.is_error]
    draft = store.load_draft(store.get_book(server.get_settings(), proj))
    assert sorted(s["title"] for s in draft["sections"]) == sorted(f"Part {i}" for i in range(8))


async def test_delete_project_removes_its_files_but_not_the_inbox(settings, model_file):
    from analyst_mcp import store

    upload_id, _ = store.save_upload(settings, model_file.read_bytes())
    assert 'to new project "Plot"' in text(await server.open_file(upload_id=upload_id, project="Plot"))
    await build_draft("Plot", SAMPLE_SPEC)  # the project's one file needs no `file`
    assert not (await server.delete_project("plot")).is_error
    assert not (settings.data_dir / "projects" / "plot").exists()
    assert not list((settings.data_dir / "uploads").iterdir())
    assert model_file.exists()
    assert error(await server.read_sheet("Plot", "Inputs"))["code"] == "project_not_found"
    assert error(await server.open_file(upload_id=upload_id))["code"] == "upload_not_found"
    # Starting it again from the inbox gives a fresh project with no draft.
    assert not (await server.open_file(inbox_file=model_file.name, project="Plot")).is_error
    assert error(await server.export_sheet("Plot"))["code"] == "no_draft"


async def test_export_is_a_copy_with_live_formulas(proj, model_file):
    before = model_file.read_bytes()
    await build_draft(proj, SAMPLE_SPEC)
    res = await server.export_sheet(proj, file_name="Client pack")
    assert not res.is_error, text(res)
    (data,) = blobs(res, "resource")
    uri = next(c.resource.uri for c in res.content if c.type == "resource")
    assert uri.startswith("file:///Client%20pack")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Dashboard", "Inputs", "Model Sheet", "Tenants", "Checks"]
    dash = wb["Dashboard"]
    cells = {c.value for row in dash.iter_rows() for c in row if isinstance(c.value, str)}
    assert "=Inputs!B1" in cells
    assert "='Model Sheet'!B2" in cells
    assert any(v.startswith("=C") and "-C" in v for v in cells)  # [r1c2]-[r2c2] became addresses
    assert len(dash._charts) == 1
    assert wb["Checks"]["B4"].value.startswith("=Dashboard!")
    assert model_file.read_bytes() == before  # the original is untouched


async def test_csv_upload_flow(settings):
    from starlette.testclient import TestClient

    client = TestClient(server.build_http_app())
    csv = b"Region,Sales\nNorth,120\nSouth,80\nNorth,40\n"
    assert client.post("/uploads", content=csv).status_code == 401
    res = client.post("/uploads", content=csv, headers={"authorization": "Bearer test-token"})
    assert res.status_code == 200, res.text
    upload_id = res.json()["upload_id"]
    opened = await server.open_file(upload_id=upload_id, file_name="sales.csv")
    assert 'to new project "sales"' in text(opened)
    q = text(await server.query("sales", 'SELECT Region, SUM(Sales) AS total FROM "sales" GROUP BY 1 ORDER BY 1'))
    assert "| North | 160 |" in q
    old_xls = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest"
    bad = client.post("/uploads", content=old_xls, headers={"authorization": "Bearer test-token"})
    assert bad.status_code == 400 and "xls" in bad.json()["error"]["message"]


def test_run_python_is_off_by_default():
    assert "run_python" not in {t.name for t in server.mcp._tool_manager.list_tools()}


async def test_pyrun_copies_input_and_collects_outputs(proj, settings):
    from analyst_mcp import pyrun, store

    files = store.get_book(settings, proj)
    code = (
        "import matplotlib.pyplot as plt\n"
        "df = pd.read_excel(INPUT, sheet_name='Tenants')\n"
        "print(len(df))\n"
        "df.to_csv(OUT + '/t.csv', index=False)\n"
        "plt.plot([1, 2]); plt.savefig(OUT + '/c.png')\n"
    )
    res = pyrun.run(files, code, timeout=60)
    assert res.exit_code == 0, res.output
    assert res.output.splitlines()[0] == "6"  # the SUM row has no saved result, so pandas sees it as empty
    assert [p.name for p in res.images] == ["c.png"] and [p.name for p in res.files] == ["t.csv"]
    slow = pyrun.run(files, "import time; time.sleep(5)", timeout=1)
    assert slow.exit_code is None and "stopped" in slow.output


async def test_row_formats_override_columns(proj):
    spec = {
        "title": "Mixed units",
        "sections": [
            {
                "title": "Sale rate",
                "columns": ["Item", "Value", "Unit"],
                "formats": [None, "int", None],
                "rows": [
                    {"cells": ["Revenue", "='Model Sheet'!B2", "₹"], "format": "inr"},
                    ["Area", "='Model Sheet'!B1", "sq ft"],
                    {"cells": ["Margin", "='Model Sheet'!B6/'Model Sheet'!B2", ""], "formats": [None, "pct"]},
                ],
            }
        ],
    }
    body = text((await build_draft(proj, spec))[-1])
    assert "| Revenue | ₹ 58,12,56,000 | ₹ |" in body
    assert "| Area | 19,375 | sq ft |" in body
    assert "| Margin | 85.8% |  |" in body
    res = await server.export_sheet(proj)
    (data,) = blobs(res, "resource")
    ws = openpyxl.load_workbook(io.BytesIO(data))["Dashboard"]
    formats = [ws.cell(r, 3).number_format for r in range(6, 9)]  # title row 4, header 5, data 6-8
    assert formats[1] == "#,##0" and formats[2] == "0.0%" and "₹" in formats[0]


async def test_breakdown_lists_the_parts_of_a_total(proj):
    body = text(await server.breakdown(proj, "'Model Sheet'!B6"))
    # Net profit = Revenue - (Construction + Approvals): the total cost subtotal is opened, signs kept.
    assert "| B2 | Revenue |" in body
    assert "| B3 | Construction | minus" in body
    assert "| B4 | Approvals | minus" in body
    assert "B5" not in body.split("items:")[1]


async def test_breakdown_groups_must_cover_every_item_once(proj):
    missing = error(await server.breakdown(proj, "'Model Sheet'!B5", {"Building": ["B3"]}))
    assert missing["code"] == "bad_groups" and "B4 Approvals" in missing["message"]
    twice = error(await server.breakdown(proj, "'Model Sheet'!B5", {"A": ["B3:B4"], "B": ["B4"]}))
    assert "in both 'A' and 'B'" in twice["message"]


async def test_sections_build_a_draft_that_exports(proj):
    area = {"title": "1. Area", "columns": ["Item", "Sq. ft."], "formats": [None, "int"],
            "rows": [["Saleable area", "='Model Sheet'!B1"]]}  # fmt: skip
    first = await server.add_section(proj, area, title="Project dashboard")
    assert len(first.content) == 2 and "Draft sections now: 1. Area." in text(first)

    body = text(await server.breakdown(proj, "'Model Sheet'!B5", {"Building": ["B3"], "Approvals": ["B4"]}))
    section = json.loads(body.split("as the user asked):\n", 1)[1].split("\n\n", 1)[0])
    check = json.loads(body.split("Its check:\n", 1)[1].split("\n\n", 1)[0])
    section["title"] = "2. Where the cost goes"
    second = await server.add_section(proj, section, checks=[check])
    assert "- OK: Heads add up to Total cost" in text(second)
    assert "| Total |" in text(second) and "100.0%" in text(second)
    assert len(blobs(second, "image")) == 2  # the section and its checks, not section 1 again

    # Change section 1 in place; the check of section 2 still points at section 2.
    area["rows"][0][0] = "Saleable carpet area"
    assert "1. Area | 2. Where the cost goes" in text(await server.add_section(proj, area, number=1))
    assert "1. Area; 2. Where the cost goes" in text(await server.list_projects())
    removed = await server.add_section(proj, number=1, remove=True)
    assert "Draft sections now: 2. Where the cost goes." in text(removed)
    await server.add_section(proj, area)  # now section 2; the cost check moved to section 1 with its section

    exported = await server.export_sheet(proj)
    (data,) = blobs(exported, "resource")
    book = openpyxl.load_workbook(io.BytesIO(data))
    ws = book["Dashboard"]
    assert ws["B2"].value == "Project dashboard"
    assert "Saleable carpet area" in [c.value for c in ws["B"]]
    assert "- 2. Where the cost goes (3 rows)\n- 1. Area (1 rows)" in text(exported)
    assert "Heads add up to Total cost" in [row[0].value for row in book["Checks"].iter_rows(max_col=1)]


async def test_unit_slips_are_corrected_and_references_keep_their_meaning(proj):
    revenue = 1000 * 2 * 10.764 * 0.9 * 30000
    section = {
        "title": "1. Profit", "columns": ["Item", "Amount", "Share"], "formats": [None, "inr_cr", "pct"],
        "rows": [["Revenue", "='Model Sheet'!B2", ""], ["Profit", "='Model Sheet'!B6", "=[r2c2]/[r1c2]*100"],
                 {"cells": ["Total", "=SUM_ABOVE()", ""], "style": "total"}],
    }  # fmt: skip
    checks = [{"label": "Revenue matches", "left": "=[r1c2]", "right": "='Model Sheet'!B2", "tolerance": 1}]
    (res,) = await build_draft(proj, {"title": "T", "sections": [section], "checks": checks})
    body = text(res)
    assert not res.is_error, body
    assert "Units corrected in this section" in body and "row 1 column 2 was in rupees" in body
    assert "row 2 column 3 was a whole-number percent" in body
    assert f"₹ {revenue / 1e7:,.2f} Cr" in body
    assert f"{PROFIT / revenue:.1%}" in body
    assert "OK: Revenue matches" in body  # the check still compares rupees with rupees
    exported = await server.export_sheet(proj)
    assert not exported.is_error, text(exported)
    (data,) = blobs(exported, "resource")
    ws = openpyxl.load_workbook(io.BytesIO(data))["Dashboard"]
    assert any("/10^7" in str(c.value) for row in ws.iter_rows() for c in row)


async def test_preview_refuses_numbers_that_read_wrong(proj):
    def spec(fmt, value, chart=None):
        return {"title": "T", "sections": [{"title": "1. S", "columns": ["Item", "Value"], "formats": [None, fmt],
                                            "rows": [["Revenue", value]], "chart": chart}]}  # fmt: skip

    assert not (await build_draft(proj, spec("inr_cr", "='Model Sheet'!B2/10^7")))[-1].is_error
    pct = error((await build_draft(proj, spec("pct", "=326.58")))[-1])
    assert "percents are fractions" in pct["message"]
    typed = error((await build_draft(proj, spec("int", "₹100")))[-1])
    assert "typed text '₹100'" in typed["message"]
    rows = [["Revenue", "='Model Sheet'!B2", "₹"], ["Saleable area", "='Model Sheet'!B1", "sq ft"]]
    area = {"title": "T", "sections": [{"title": "1. S", "columns": ["Item", "Value", "Unit"], "formats": [None, "inr"],
                                        "rows": rows}]}  # fmt: skip
    assert "row 2 shows 19,375.20 as money" in error((await build_draft(proj, area))[-1])["message"]
    pie = {
        "title": "T",
        "sections": [{
            "title": "1. Cost", "columns": ["Head", "Amount"],
            "rows": [["Construction", "='Model Sheet'!B3"],
                     {"cells": ["Total", "='Model Sheet'!B5"], "style": "total"}],
            "chart": {"type": "pie", "value_columns": [2]},
        }],
    }  # fmt: skip
    assert "pie slices add up to" in error((await build_draft(proj, pie))[-1])["message"]


def test_big_numbers_get_words():
    from analyst_mcp.fmt import approx

    assert approx(643834212.77, "Rs.") == "≈ ₹ 64.38 Cr"
    assert approx(475000, "Rs.") == "≈ ₹ 4.75 lakh"
    assert approx(35000, "Rs.") == ""
    assert approx(643834212.77, "") == "≈ 643.83 million"
    assert approx(19375, "sq ft") == ""


async def test_a_lost_name_means_the_only_project(proj):
    # What a model that lost track sends: a chat attachment's id, the file name, nothing.
    for guess in ("863bdbbd-794", "model.xlsx", "MODEL", ""):
        res = await server.read_sheet(guess, "Inputs", "A1:B1")
        assert not res.is_error, text(res)
        assert "Plot Area" in text(res)


async def test_each_chat_keeps_to_its_project_by_name(settings, proj):
    """Two projects open at once: each name reaches its own file, a near name finds its project, and a call that
    names none is refused with the list so the model asks the user."""
    other = openpyxl.Workbook()
    other.active.title = "Meals"
    other.active.append(["Day", "kcal"])
    other.active.append(["Mon", 1800])
    other.save(settings.inbox() / "diet.xlsx")
    assert not (await server.open_file(inbox_file="diet.xlsx", project="Diet plan")).is_error
    assert "Plot Area" in text(await server.read_sheet("model", "Inputs", "A1:B1"))
    assert "1,800" in text(await server.read_sheet("diet", "Meals"))
    lost = error(await server.breakdown("", "'Model Sheet'!B5"))
    assert lost["code"] == "project_not_found" and '"Diet plan"' in lost["message"] and '"model"' in lost["message"]
    assert error(await server.find("Budget", "profit"))["code"] == "project_not_found"
    listed = text(await server.list_projects())
    assert "| Diet plan |" in listed and "| model |" in listed


async def test_switch_shows_the_project_and_never_guesses(settings, proj):
    res = await server.switch_to_project("MODEL")
    assert 'Now on project "model"' in text(res) and "| B6 | Net profit |" in text(res)
    other = openpyxl.Workbook()
    other.active["A1"] = "x"
    other.save(settings.inbox() / "b.xlsx")
    await server.open_file(inbox_file="b.xlsx", project="Budget 2026")
    assert error(await server.switch_to_project("Sales"))["code"] == "project_not_found"


def month(settings, name: str, qty: int, sheet: str = "Sales") -> None:
    wb = openpyxl.Workbook()
    wb.active.title = sheet
    wb.active.append(["Item", "Value"])
    for item, share in (("Bolt", 5), ("Nut", 3), ("Washer", 2)):
        wb.active.append([item, qty * share])
    wb.save(settings.inbox() / f"{name}.xlsx")


async def test_a_project_holds_files_and_query_reads_them_all(settings):
    """May, June and July in one project: SQL stacks their shared sheet with a `file` column, so totals across the
    months (an ABC analysis) are one query; one file's sheet is "<file>.<sheet>"."""
    for name, qty in (("may-data", 10), ("june-data", 20), ("july-data", 30)):
        month(settings, name, qty)
        res = await server.open_file(inbox_file=f"{name}.xlsx", project="Dialysis")
        assert not res.is_error, text(res)
    assert "Files in this project: 'may-data', 'june-data', 'july-data'" in text(res)
    q = text(await server.query("dialysis", 'SELECT Item, SUM(Value) AS v FROM "Sales" GROUP BY 1 ORDER BY v DESC'))
    assert "| Bolt | 300 |" in q and "| Washer | 120 |" in q
    june = text(await server.query("Dialysis", 'SELECT SUM(Value) AS v FROM "june-data.Sales"'))
    assert "| 200 |" in june
    files = text(await server.query("Dialysis", "SELECT DISTINCT file FROM cells ORDER BY 1"))
    assert "| july-data |" in files and "| may-data |" in files
    found = text(await server.find("Dialysis", "washer"))
    assert "[may-data] Sales!A4" in found and "[july-data] Sales!A4" in found
    listed = text(await server.list_projects())
    assert "| Dialysis | may-data, june-data, july-data |" in listed


async def test_one_file_tools_ask_which_file_unless_the_sheet_says(settings):
    month(settings, "rates", 1, sheet="Rates")
    month(settings, "stock", 2, sheet="Stock")
    for name in ("rates", "stock"):
        await server.open_file(inbox_file=f"{name}.xlsx", project="Dialysis")
    assert "Bolt" in text(await server.read_sheet("Dialysis", "Rates"))  # only one file has that sheet
    month(settings, "stock2", 3, sheet="Rates")
    await server.open_file(inbox_file="stock2.xlsx", project="Dialysis")
    err = error(await server.read_sheet("Dialysis", "Rates"))
    assert err["code"] == "file_needed" and '"rates"' in err["message"] and '"stock2"' in err["message"]
    assert "Bolt" in text(await server.read_sheet("Dialysis", "Rates", file="stock2"))


async def test_the_same_name_is_a_new_version_and_files_move(settings):
    month(settings, "june", 1)
    await server.open_file(inbox_file="june.xlsx", project="Q2")
    month(settings, "june", 2)
    again = text(await server.open_file(inbox_file="june.xlsx", project="Q2"))
    assert 'Saved the new "june" as version 2' in again
    assert "| 10 |" in text(await server.query("Q2", 'SELECT Value FROM "Sales" WHERE Item = \'Bolt\''))
    month(settings, "rates", 1)
    await server.open_file(inbox_file="rates.xlsx")  # a project of its own, named after the file
    moved = text(await server.move_file("rates", "rates", "Q2"))
    assert 'now in project "Q2"' in moved
    listed = text(await server.list_projects())
    assert "| Q2 | june (1 edits), rates |" in listed and "| rates |" not in listed


async def test_projects_from_before_files_still_open(settings):
    """A project saved as projects/<p>/versions/ (one file per project) becomes a project with that one file."""
    folder = settings.data_dir / "projects" / "old-plan"
    (folder / "versions").mkdir(parents=True)
    month(settings, "tmp", 1)
    (folder / "versions" / "1.xlsx").write_bytes((settings.inbox() / "tmp.xlsx").read_bytes())
    (folder / "project.json").write_text('{"name": "Old plan", "version": 1, "updated_at": "2026-10-09"}')
    assert "Bolt" in text(await server.read_sheet("Old plan", "Sales"))
    assert (folder / "files" / "old-plan" / "versions" / "1.xlsx").exists()


async def test_calls_at_once_migrate_an_old_project_once(settings):
    """Tools run in threads; two reads of a project saved the old way must not both move its folder."""
    import asyncio

    folder = settings.data_dir / "projects" / "old-plan"
    (folder / "versions").mkdir(parents=True)
    month(settings, "tmp", 1)
    (folder / "versions" / "1.xlsx").write_bytes((settings.inbox() / "tmp.xlsx").read_bytes())
    (folder / "project.json").write_text('{"name": "Old plan", "version": 1}')
    results = await asyncio.gather(*[server.read_sheet("Old plan", "Sales") for _ in range(4)])
    assert all(not r.is_error for r in results), [text(r) for r in results]
